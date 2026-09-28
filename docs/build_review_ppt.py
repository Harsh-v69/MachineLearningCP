"""Builds docs/CP_Project_Review.pptx: a plain, fully editable review deck.

Everything is native PowerPoint text, tables and shapes (no images), so it
opens and edits cleanly in PowerPoint and Canva. Numbers are copied from
results/*.md; if results change, update the tables here and rerun:
    python docs/build_review_ppt.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

INK, MUTED, ACCENT, RULE = RGBColor(0x1A, 0x1A, 0x1A), RGBColor(0x5A, 0x5A, 0x5A), RGBColor(0x1F, 0x4F, 0xC8), RGBColor(0xC8, 0xC8, 0xC8)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, content, size=20, bold=False, color=INK, font=FONT, align=None, anchor=None):
    """content: str or list of str / (str, dict) paragraphs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    items = content if isinstance(content, list) else [content]
    for i, it in enumerate(items):
        t, o = (it, {}) if isinstance(it, str) else it
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = t
        r.font.name = o.get("font", font)
        r.font.size = Pt(o.get("size", size))
        r.font.bold = o.get("bold", bold)
        r.font.color.rgb = o.get("color", color)
        p.space_after = Pt(o.get("after", 6))
        if align:
            p.alignment = align
    return tb


def bullets(slide, x, y, w, h, items, size=20):
    return text(slide, x, y, w, h, ["•  " + i if isinstance(i, str) else i for i in items], size=size)


def new_slide(title, subtitle=None):
    s = prs.slides.add_slide(BLANK)
    text(s, 0.6, 0.35, 12, 0.8, title, size=34, bold=True)
    line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.6), Inches(1.15), Inches(12.7), Inches(1.15))
    line.line.color.rgb = ACCENT
    line.line.width = Pt(2)
    if subtitle:
        text(s, 0.6, 1.22, 12, 0.5, subtitle, size=16, color=MUTED)
    return s


def box(slide, x, y, w, h, label, sub=None, fill=None, size=16):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill or RGBColor(0xF2, 0xF5, 0xFC)
    sh.line.color.rgb = ACCENT
    sh.line.width = Pt(1.5)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), True, INK
    if sub:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = sub
        r2.font.name, r2.font.size, r2.font.color.rgb = FONT, Pt(size - 4), MUTED
    return sh


def arrow(slide, x1, y1, x2, y2):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = INK
    c.line.width = Pt(2)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", {"type": "triangle"})
    ln.append(tail)
    return c


def table(slide, x, y, w, rows, col_w=None, size=14, row_h=0.42):
    t = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows))).table
    for ci, cw in enumerate(col_w or []):
        t.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.name, r.font.size = FONT, Pt(size)
                    r.font.bold = ri == 0
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if ri == 0 else INK
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT if ri == 0 else (RGBColor(0xF6, 0xF6, 0xF6) if ri % 2 == 0 else RGBColor(0xFF, 0xFF, 0xFF))
    return t


# 1. Title -------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
text(s, 0.8, 0.9, 11.7, 0.5, "CP PROJECT REVIEW", size=16, bold=True, color=ACCENT)
text(s, 0.8, 1.4, 11.7, 2.0, "Adaptive Graph Validation and Solver Selection for Plan-Graph Planning", size=40, bold=True)
text(s, 0.8, 3.2, 11.7, 0.6, "An LLM-free extension of TAPE (Tool-Guided Adaptive Planning and Constrained Execution)", size=20, color=MUTED)
text(s, 0.8, 4.3, 6.0, 2.6, [
    ("Subject: Machine Learning", {"bold": True}), "Guide: Dr. Premanand P Ghadekar", "Group 3, Year 3, Semester 1"], size=20)
text(s, 7.0, 4.3, 5.6, 2.6, [("Group members", {"bold": True}), "Divyanshu Gajbhiye", "Vidhi Hadoltikar", "Harshvardhan Rawat", "Nirmit Hatti"], size=20)

# 2. Introduction ------------------------------------------------------
s = new_slide("Introduction")
bullets(s, 0.6, 1.5, 12.1, 5.5, [
    "Long-horizon planning is hard for language-model agents: a single plan often fails midway.",
    "TAPE (Jeong et al., 2026) generates several candidate plans, merges them into one plan graph, picks a path with an external solver (ILP), executes it, and replans when reality differs.",
    "A plan graph is only as good as the plans behind it, and the solver is fixed in advance.",
    "Our project adds a checking and selection layer on top of the TAPE pipeline: validate moves, learn from failures, score moves, and choose the solver per task.",
    "Tested on four benchmark families: Sokoban, River Crossing, Rush Hour and Boxoban.",
], size=22)

# 3. Motivation --------------------------------------------------------
s = new_slide("Motivation")
bullets(s, 0.6, 1.5, 6.1, 5.5, [
    "Nothing in TAPE checks a proposed move against the real rules before it is trusted.",
    "Errors are found only after execution, as a mismatch.",
    "One solver is used for every task, whether or not it is needed.",
    "Past failures are forgotten between episodes.",
], size=22)
text(s, 7.0, 1.5, 5.7, 0.5, "What goes wrong without a checker", size=20, bold=True, color=ACCENT)
bullets(s, 7.0, 2.1, 5.7, 4.5, [
    "River Crossing: a plan can reach the goal while breaking the safety rule on the way.",
    "Sokoban: a box pushed into a corner can never be recovered.",
    "Rush Hour: two blocked vehicles can seal a row.",
    "Our measurements show the plain pipeline reaching the River Crossing goal every time, but never safely.",
], size=20)

# 4. Literature survey -------------------------------------------------
s = new_slide("Literature Survey")
table(s, 0.6, 1.45, 12.1, [
    ["Work", "Idea", "Gap for our project"],
    ["ReAct (Yao et al., 2022)", "Interleave reasoning and actions with an environment", "No plan-level check before acting"],
    ["Tree of Thoughts (Yao et al., 2023)", "Search over a tree of reasoning steps with backtracking", "Tree, not a merged graph"],
    ["Graph of Thoughts (Besta et al., 2023)", "Thoughts as a graph; paths merge and reuse", "No environment-rule validation"],
    ["Reflexion (Shinn et al., 2023)", "Learn from failures through verbal feedback", "Feedback is text, not a per-move confidence"],
    ["Plan-and-Solve, ReWOO (2023)", "Separate planning from execution", "Single plan, no solver"],
    ["TAPE (Jeong et al., 2026)", "Plan graph, external solver, constrained execution, replanning", "Unchecked graph; fixed solver"],
    ["A* (Hart et al., 1968)", "Heuristic shortest-path search", "Used here as the alternative selector"],
    ["Sokoban / Rush Hour complexity (Culberson 1997; Flake and Baum 2002)", "Both puzzles are PSPACE-complete", "Hard, structured test beds"],
], col_w=[3.6, 4.7, 3.8], size=13, row_h=0.55)

# 5. Objectives --------------------------------------------------------
s = new_slide("Objectives")
bullets(s, 0.6, 1.5, 12.1, 5.6, [
    "Reproduce the core TAPE pipeline as a baseline: plans, plan graph, solver, constrained execution, replanning.",
    "Validate states and transitions against environment rules before trusting them: reject the impossible, downweight the uncertain.",
    "Use execution feedback to raise or lower confidence in later episodes.",
    "Define a dynamic score for each move from goal progress, confidence and remaining budget.",
    "Choose between A* and a constraint solver per task instead of always using one.",
    "Compare every extension with the baseline through controlled experiments and ablations.",
], size=22)

# 6. Project flow diagram ----------------------------------------------
s = new_slide("Project Flow Diagram")
BW, STEP, Y = 1.6, 1.85, 2.5
box(s, 0.3, Y, 1.6, 1.2, "Environment", "state and rules", fill=RGBColor(0xEE, 0xEE, 0xEE), size=14)
labels = [("Propose", "candidate plans"), ("Merge", "plan graph"), ("Check", "validator + memory"), ("Choose", "score + A* / CP-SAT"), ("Run", "execute path")]
x = 2.3
for i, (a, b) in enumerate(labels):
    box(s, x, Y, BW, 1.2, a, b, size=15)
    arrow(s, x - (STEP - BW) if i else 1.9, Y + 0.6, x, Y + 0.6)
    x += STEP
arrow(s, x - (STEP - BW), Y + 0.6, x, Y + 0.6)
box(s, x, Y, 1.3, 1.2, "Validated path", None, fill=RGBColor(0xEE, 0xEE, 0xEE), size=14)
run_cx, prop_cx = 2.3 + 4 * STEP + BW / 2, 2.3 + BW / 2
fb = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(run_cx), Inches(Y + 1.2), Inches(run_cx), Inches(Y + 2.1))
fb.line.color.rgb, fb.line.width = INK, Pt(2)
fb2 = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(run_cx), Inches(Y + 2.1), Inches(prop_cx), Inches(Y + 2.1))
fb2.line.color.rgb, fb2.line.width = INK, Pt(2)
arrow(s, prop_cx, Y + 2.1, prop_cx, Y + 1.2)
text(s, prop_cx + 0.3, Y + 2.2, 7.0, 0.5, "Mismatch: record the failure, replan from the real state", size=16, color=MUTED)
text(s, 0.6, 5.7, 12.1, 1.4, [
    ("Phase 1: Propose, Merge, Run.  Phase 2 and 3: Check (rules, then memory of past failures).  Phase 4 and 5: Choose (move score, then A* or CP-SAT).", {"size": 16, "color": MUTED}),
    ("Proposers are synthetic programs; no language model is used.", {"size": 16, "color": MUTED})])

# 7. Algorithm ---------------------------------------------------------
s = new_slide("Algorithm")
text(s, 0.6, 1.45, 7.4, 5.4, [
    ("for each episode, from the current real state:", {"font": "Consolas", "size": 16, "bold": True}),
    ("  plans  = proposer(state)", {"font": "Consolas", "size": 16}),
    ("  graph  = merge(plans)       # equal states share a node", {"font": "Consolas", "size": 16}),
    ("  for each move in graph:", {"font": "Consolas", "size": 16}),
    ("      v    = validator(state, move)   # reject | accept(conf)", {"font": "Consolas", "size": 16}),
    ("      conf = v.conf * experience(state, move)", {"font": "Consolas", "size": 16}),
    ("      cost = score(progress, conf, budget)", {"font": "Consolas", "size": 16}),
    ("  path = A* if one piece moves else CP-SAT", {"font": "Consolas", "size": 16}),
    ("  execute(path)", {"font": "Consolas", "size": 16}),
    ("  on mismatch: record failure, replan", {"font": "Consolas", "size": 16}),
], size=17)
text(s, 8.3, 1.45, 4.5, 0.5, "Key formulas", size=20, bold=True, color=ACCENT)
text(s, 8.3, 2.0, 4.5, 4.8, [
    ("Move cost", {"bold": True, "size": 16}),
    ("cost = step + 1.5 x regression + 4.0 x (1 - confidence) + 2.0 x (steps used / budget)^2", {"size": 15}),
    ("Experience confidence", {"bold": True, "size": 16}),
    ("(successes + 1) / (successes + failures + 2)", {"size": 15}),
    ("Solver rule", {"bold": True, "size": 16}),
    ("one movable piece: A*; several pieces: CP-SAT", {"size": 15}),
])

# 8. Coding progress ---------------------------------------------------
s = new_slide("Coding Progress", "Target for this review: 40%. Completed: all five pipeline phases and the evaluation experiments.")
table(s, 0.6, 1.9, 12.1, [
    ["Phase", "What was coded", "Status"],
    ["1  TAPE baseline", "Environment interface, plan graph, CP-SAT solver, constrained execution, replanning", "Done"],
    ["2  Graph validation", "Rule checkers: Sokoban corners, River Crossing safety, Rush Hour gridlock", "Done"],
    ["3  Self-improving graph", "SQLite experience store with smoothed confidence", "Done"],
    ["4  Dynamic score", "Per-move cost from progress, confidence and budget", "Done"],
    ["5  Adaptive selection", "A* and CP-SAT, chosen by task complexity", "Done"],
    ["6  Evaluation", "Ablation, memory test, weight tuning, proposer sweep, Boxoban", "Core done"],
], col_w=[3.0, 7.3, 1.8], size=15, row_h=0.55)
text(s, 0.6, 6.0, 12.1, 1.0, "About 1,600 lines in src/tape, 1,050 in experiments, 750 in tests. 66 automated tests pass. Code: github.com/Harsh-v69/MachineLearningCP", size=16, color=MUTED)

# 9. Coding: key parts -------------------------------------------------
s = new_slide("Coding: Key Components")
text(s, 0.6, 1.45, 6.0, 0.5, "Validator, Sokoban (simplified)", size=18, bold=True, color=ACCENT)
text(s, 0.6, 1.95, 6.0, 4.6, [
    ("for box in moved_boxes:", {"font": "Consolas", "size": 14}),
    ("    if box in goals: continue", {"font": "Consolas", "size": 14}),
    ("    if is_corner(box):", {"font": "Consolas", "size": 14}),
    ("        return reject(\"corner deadlock\")", {"font": "Consolas", "size": 14}),
    ("    if against_wall(box):", {"font": "Consolas", "size": 14}),
    ("        return accept(conf=0.5)", {"font": "Consolas", "size": 14}),
    ("return accept(conf=1.0)", {"font": "Consolas", "size": 14}),
    ("", {"size": 8}),
    ("Rejected moves never enter the plan graph. Doubtful moves stay in, but cost more.", {"size": 16, "color": MUTED}),
])
text(s, 6.9, 1.45, 5.8, 0.5, "Adaptive solver choice (simplified)", size=18, bold=True, color=ACCENT)
text(s, 6.9, 1.95, 5.8, 4.6, [
    ("def decide_method(level, graph):", {"font": "Consolas", "size": 14}),
    ("    if level.complexity(graph.start) > 1:", {"font": "Consolas", "size": 14}),
    ("        return CP_SAT", {"font": "Consolas", "size": 14}),
    ("    return ASTAR", {"font": "Consolas", "size": 14}),
    ("", {"size": 8}),
    ("Every game implements one interface: ACTIONS, step, is_goal, heuristic, complexity, render. A new game plugs in without touching the pipeline.", {"size": 16, "color": MUTED}),
    ("", {"size": 8}),
    ("Benchmarks: Sokoban (4 levels), River Crossing, Rush Hour, Boxoban (3,000 real levels).", {"size": 16, "color": MUTED}),
])

# 10. Results 1 --------------------------------------------------------
s = new_slide("Results and Discussion: Safety and Speed", "Ablation, 30 episodes per cell, 15% slip, synthetic proposer")
table(s, 0.6, 1.9, 5.8, [
    ["River Crossing", "Safe success"],
    ["Baseline (no checker)", "0%"],
    ["With validator", "100%"],
    ["Goal reached, baseline", "100% (unsafely)"],
], col_w=[3.6, 2.2], size=15, row_h=0.55)
table(s, 6.8, 1.9, 5.9, [
    ["Planning time (ms)", "Always CP-SAT", "Adaptive"],
    ["Sokoban, 1 box", "51.8", "7.6"],
    ["River Crossing", "62.1", "3.6"],
    ["Sokoban, 2 boxes", "99.9", "96.4"],
    ["Rush Hour", "36.2", "36.1"],
], col_w=[2.7, 1.7, 1.5], size=15, row_h=0.5)
bullets(s, 0.6, 4.6, 12.1, 2.4, [
    "Safe success means the goal was reached and no move broke a rule, judged by an independent auditor.",
    "The validator turns unsafe successes into safe ones. The adaptive solver saves time where one piece moves and ties elsewhere.",
    "Hard Sokoban stays near 53% in every configuration: our extensions do not change it.",
], size=18)

# 11. Results 2: Boxoban ----------------------------------------------
s = new_slide("Results and Discussion: Boxoban", "Share of 1,000 real levels per set solved within a node budget: no pruning vs with checker")
table(s, 0.6, 1.95, 12.1, [
    ["Nodes per search", "Unfiltered", "Medium", "Hard"],
    ["2,000", "57% vs 65%", "14% vs 23%", "6% vs 13%"],
    ["5,000", "73% vs 82%", "29% vs 45%", "16% vs 29%"],
    ["10,000", "84% vs 92%", "46% vs 65%", "27% vs 49%"],
    ["20,000", "92% vs 96%", "61% vs 82%", "44% vs 70%"],
    ["30,000", "95% vs 98%", "71% vs 88%", "54% vs 82%"],
], col_w=[3.1, 3.0, 3.0, 3.0], size=16, row_h=0.5)
bullets(s, 0.6, 5.2, 12.1, 2.0, [
    "The checker wins in all 15 cells, with non-overlapping 95% intervals. It never missed a level the other search solved.",
    "It also expands 20% to 25% fewer nodes. Pruning dead corners is standard in Sokoban solvers; this shows our checker works as one on real levels.",
], size=18)

# 12. Results 3: memory, tuning, limits --------------------------------
s = new_slide("Results and Discussion: Memory, Tuning, Limits")
table(s, 0.6, 1.5, 6.4, [
    ["Experience store", "Change in replans"],
    ["Sokoban, 1 box", "-0.56 +/- 0.18"],
    ["River Crossing", "-0.54 +/- 0.15"],
    ["Rush Hour", "-0.17 +/- 0.11"],
    ["Sokoban, 2 boxes", "+0.06 +/- 0.03"],
], col_w=[3.6, 2.8], size=15, row_h=0.5)
text(s, 0.6, 4.2, 6.4, 0.9, "Memory helps when failures depend on the state and the candidates contain a detour. Negative is better.", size=15, color=MUTED)
text(s, 7.3, 1.5, 5.4, 0.5, "What the data does not show", size=20, bold=True, color=ACCENT)
bullets(s, 7.3, 2.1, 5.4, 4.8, [
    "Proposers are synthetic: no language model was used.",
    "A checker filters plans; it cannot invent one.",
    "Tuning the cost weights (80-point grid, held-out check) found nothing better than the defaults.",
    "Merging plans did not shorten paths: cost, not move count, is minimized.",
    "With a decent proposer, goal-reach rate barely changes; the value is safety and fewer wasted replans.",
], size=17)

# 13. Conclusion -------------------------------------------------------
s = new_slide("Conclusion")
bullets(s, 0.6, 1.5, 12.1, 3.4, [
    "We built a checking and selection layer on the TAPE pipeline and tested it on four benchmark families.",
    "Rule checking makes plans safe (River Crossing 0% to 100% safe) and makes search solve more real levels (Boxoban, up to 28 points).",
    "Choosing A* or CP-SAT per task cuts planning time where one piece moves.",
    "Remembering failures reduces replans, but only when failures are state-dependent.",
], size=21)
text(s, 0.6, 5.0, 12.1, 0.5, "Future work", size=20, bold=True, color=ACCENT)
bullets(s, 0.6, 5.5, 12.1, 1.6, [
    "A learned, non-LLM proposer with realistic errors; more Boxoban files and a second checker rule.",
    "Write up the results as a paper (draft in progress).",
], size=18)

# 14. References -------------------------------------------------------
refs = [
    "[1] J. Jeong, J. Kim, K. Lee, \"TAPE: Tool-Guided Adaptive Planning and Constrained Execution in Language Model Agents,\" arXiv:2602.19633, 2026.",
    "[2] S. Yao et al., \"ReAct: Synergizing Reasoning and Acting in Language Models,\" arXiv:2210.03629, 2022.",
    "[3] S. Yao et al., \"Tree of Thoughts: Deliberate Problem Solving with Large Language Models,\" arXiv:2305.10601, 2023.",
    "[4] M. Besta et al., \"Graph of Thoughts: Solving Elaborate Problems with Large Language Models,\" arXiv:2308.09687, 2023.",
    "[5] N. Shinn et al., \"Reflexion: Language Agents with Verbal Reinforcement Learning,\" arXiv:2303.11366, 2023.",
    "[6] T. Schick et al., \"Toolformer: Language Models Can Teach Themselves to Use Tools,\" arXiv:2302.04761, 2023.",
    "[7] J. Wei et al., \"Chain-of-Thought Prompting Elicits Reasoning in Large Language Models,\" arXiv:2201.11903, 2022.",
    "[8] L. Wang et al., \"Plan-and-Solve Prompting,\" arXiv:2305.04091, 2023.",
    "[9] B. Xu et al., \"ReWOO: Decoupling Reasoning from Observations for Efficient Augmented Language Models,\" arXiv:2305.18323, 2023.",
    "[10] M. Shridhar et al., \"ALFWorld: Aligning Text and Embodied Environments for Interactive Learning,\" arXiv:2010.03768, 2020.",
    "[11] H. Trivedi et al., \"MuSiQue: Multihop Questions via Single-hop Question Composition,\" arXiv:2108.00573, 2021.",
    "[12] K. Cobbe et al., \"Training Verifiers to Solve Math Word Problems,\" arXiv:2110.14168, 2021.",
    "[13] P. Hart, N. Nilsson, B. Raphael, \"A Formal Basis for the Heuristic Determination of Minimum Cost Paths,\" IEEE Trans. Systems Science and Cybernetics, 1968.",
    "[14] J. Culberson, \"Sokoban is PSPACE-complete,\" Tech. Rep. TR97-02, Univ. of Alberta, 1997.",
    "[15] G. W. Flake, E. B. Baum, \"Rush Hour is PSPACE-complete,\" Theoretical Computer Science 270(1-2):895-911, 2002.",
    "[16] DeepMind, \"Boxoban levels,\" github.com/google-deepmind/boxoban-levels.",
    "[17] L. Perron, V. Furnon, \"OR-Tools,\" Google, developers.google.com/optimization.",
]
s = new_slide("References")
text(s, 0.6, 1.4, 6.1, 5.9, [(r, {"size": 13, "after": 6}) for r in refs[:8]])
text(s, 6.9, 1.4, 5.9, 5.9, [(r, {"size": 13, "after": 6}) for r in refs[8:]])

out = Path(__file__).resolve().parent / "CP_Project_Review.pptx"
prs.save(out)
print("wrote", out)
