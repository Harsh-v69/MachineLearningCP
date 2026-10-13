"""Builds paper/paper.docx in the ICCISD two-column IEEE conference format from
paper/paper.md and paper/figures/. Layout values were measured from the
conference template PDF: A4, Times New Roman 10 pt, two columns 245.8 pt wide
with a 15.4 pt gap, small-caps centred section headings with Roman numerals,
italic lettered subheadings, a 9 pt bold abstract, bordered tables with centred
upper-case captions, 10 pt references, and the conference header and copyright
footer on the first page only.

Run: python paper/make_figures.py && python paper/build_paper_docx.py
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
FONT = "Times New Roman"
ESC = "\u0001"
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]

CONFERENCE = "2026 International Conference on Computational Intelligence Systems and Devices (ICCISD)"
COPYRIGHT = "979-8-3315-6371-4/26/$31.00 ©2026 IEEE"
AUTHORS = [  # (name, email); department and institute are the same for the group
    ("Divyanshu Gajbhiye", "[email to add]"),
    ("Vidhi Hadoltikar", "[email to add]"),
    ("Harshvardhan Rawat", "harshvardhan.rawat24@vit.edu"),
    ("Nirmit Hatti", "nirmit.hatti24@vit.edu"),
]
DEPT, INSTITUTE, CITY = "CSE (AI&ML)", "Vishwakarma Institute of Technology", "Pune, India"

# ---------------------------------------------------------------- page setup
doc = Document()
sec1 = doc.sections[0]
sec1.page_width, sec1.page_height = Pt(595.22), Pt(842)
for s in (sec1,):
    s.left_margin, s.right_margin = Pt(45.4), Pt(42.7)
    s.top_margin, s.bottom_margin = Pt(54), Pt(36)
    s.header_distance, s.footer_distance = Pt(28), Pt(24)
    s.different_first_page_header_footer = True


def fix_fonts(style):
    style.font.name = FONT
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for att in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(att), FONT)
    for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(att)) is not None:
            del rf.attrib[qn(att)]
    style.font.color.rgb = RGBColor(0, 0, 0)


normal = doc.styles["Normal"]
fix_fonts(normal)
normal.font.size = Pt(10)
pf = normal.paragraph_format
pf.space_before = pf.space_after = Pt(0)
pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for nm in ("Title", "Heading 1", "Heading 2"):
    st = doc.styles[nm]
    fix_fonts(st)
    st.font.bold = False
    st.font.italic = False
    st.font.size = Pt(10)
    st.paragraph_format.keep_with_next = True
    st.element.get_or_add_pPr()
# Title style often carries a bottom border; remove it
for pbdr in doc.styles["Title"].element.xpath(".//w:pBdr"):
    pbdr.getparent().remove(pbdr)


def runs(par, text, size=None, bold=None, italic=None, ital=False):
    """Bold is **text**. Italics (*text*) are parsed only when ital=True (the reference
    list), so that A*, AO* and LAO* in running text are left alone."""
    text = text.replace("\\*", ESC)
    pattern = r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*)" if ital else r"(\*\*.+?\*\*)"
    for tok in re.split(pattern, text):
        if not tok:
            continue
        b = tok.startswith("**") and tok.endswith("**") and len(tok) > 4
        it = ital and (not b) and tok.startswith("*") and tok.endswith("*") and len(tok) > 2
        body = tok[2:-2] if b else tok[1:-1] if it else tok
        r = par.add_run(body.replace(ESC, "*"))
        r.font.name = FONT
        if size:
            r.font.size = Pt(size)
        if b or bold:
            r.bold = True
        if it or italic:
            r.italic = True


def para(text="", align=None, indent=False, size=None, bold=None, italic=None, before=0, after=0, keep_next=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if indent:
        p.paragraph_format.first_line_indent = Pt(14.4)
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
    p.paragraph_format.keep_with_next = keep_next
    if text:
        runs(p, text, size=size, bold=bold, italic=italic)
    return p


def tiny(par):
    """Collapse an empty header or footer paragraph to almost nothing."""
    par.paragraph_format.space_before = par.paragraph_format.space_after = Pt(0)
    par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    par.paragraph_format.line_spacing = Pt(1)
    r = par.add_run("")
    r.font.size = Pt(1)


# headers and footers: the conference line and the copyright line appear on page 1 only
fh = sec1.first_page_header.paragraphs[0]
fh.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fh.add_run(CONFERENCE)
r.font.name, r.font.size = FONT, Pt(8)
ff = sec1.first_page_footer.paragraphs[0]
ff.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = ff.add_run(COPYRIGHT)
r.font.name, r.font.size = FONT, Pt(8)
tiny(sec1.header.paragraphs[0])
tiny(sec1.footer.paragraphs[0])

# ----------------------------------------------------------- title and authors
lines = (HERE / "paper.md").read_text(encoding="utf-8").splitlines()
title = next(l[2:] for l in lines if l.startswith("# "))
tp = para(align=WD_ALIGN_PARAGRAPH.CENTER, before=10, after=12)
r = tp.add_run(title)
r.bold, r.font.size, r.font.name = True, Pt(24), FONT

tbl = doc.add_table(rows=2, cols=3)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl.autofit = False
slots = [(0, 0), (0, 1), (0, 2), (1, 1)]  # four authors: three in the first row, one centred under them
ordinals = ["st", "nd", "rd", "th"]
for (name, email), (ri, ci), sup in zip(AUTHORS, slots, ordinals):
    cell = tbl.cell(ri, ci)
    cell.width = Pt(168)
    first = cell.paragraphs[0]
    first.alignment = WD_ALIGN_PARAGRAPH.CENTER
    n = AUTHORS.index((name, email)) + 1
    a = first.add_run(str(n))
    a.font.size, a.font.name = Pt(9), FONT
    b = first.add_run(sup)
    b.font.size, b.font.name, b.font.superscript = Pt(9), FONT, True
    c = first.add_run(f" {name}")
    c.font.size, c.font.name = Pt(9), FONT
    for t in (DEPT, INSTITUTE, CITY, email):
        q = cell.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.space_before = q.paragraph_format.space_after = Pt(0)
        rr = q.add_run(t)
        rr.font.size, rr.font.name = Pt(9), FONT
    for pp in cell.paragraphs:
        pp.paragraph_format.space_after = Pt(0)
for ci in range(3):
    tbl.cell(1, ci).paragraphs[0].paragraph_format.space_after = Pt(0)
# row spacing between the two author rows
tbl.rows[0].height = Pt(66)
doc.add_paragraph().paragraph_format.space_after = Pt(4)

# ------------------------------------------------------------ two-column body
sec2 = doc.add_section(WD_SECTION.CONTINUOUS)
for s in (sec2,):
    s.left_margin, s.right_margin = Pt(45.4), Pt(42.7)
    s.top_margin, s.bottom_margin = Pt(54), Pt(36)
    s.header_distance, s.footer_distance = Pt(28), Pt(24)
    s.different_first_page_header_footer = False
sec2.header.is_linked_to_previous = True
sec2.footer.is_linked_to_previous = True
cols = sec2._sectPr.xpath("./w:cols")
col = cols[0] if cols else OxmlElement("w:cols")
col.set(qn("w:num"), "2")
col.set(qn("w:space"), "308")  # 15.4 pt in twentieths of a point
if not cols:
    sec2._sectPr.append(col)


def shade_border(cell, borders=True):
    tcPr = cell._tc.get_or_add_tcPr()
    if borders:
        tb = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), "000000")
            tb.append(e)
        tcPr.append(tb)


def add_table(rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    n = len(rows[0])
    weight = [max(10, max(len(r[j].strip()) for r in rows)) for j in range(n)]
    tblPr = t._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for edge in ("left", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:w"), "60")
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.append(mar)
    widths = [Pt(245 * w / sum(weight)) for w in weight]
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.width = widths[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(1.5)
            rr = p.add_run(val.strip())
            rr.font.size, rr.font.name = Pt(9 if len(rows[0]) <= 4 else 7.5), FONT
            rr.bold = True if i == 0 else None
            shade_border(cell)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(6)
    return t


def search_rows():
    S = json.loads((ROOT / "results" / "search_methods.json").read_text(encoding="utf-8"))["games"]
    algs = [("BFS", "BFS"), ("DFS", "DFS"), ("weighted A*", "A*"), ("meet in the middle", "MITM")]
    short = {"Peg Solitaire, 22 jumps out": "Peg, 22 jumps", "24-puzzle, 20 move scramble": "24-puzzle, 20",
             "24-puzzle, 40 move scramble": "24-puzzle, 40", "Rubik's cube, 5 move scramble": "Cube, 5 moves",
             "Rubik's cube, 6 move scramble": "Cube, 6 moves", "Boxoban (20 levels)": "Boxoban", "Peg Solitaire": "Peg, real start"}
    rows = [["Game"] + [b for _, b in algs]]
    for g, v in S.items():
        cells = []
        for a, _ in algs:
            c = v.get(a)
            cells.append("n/a" if c is None else "0" if c["nodes"] is None else f"{c['solved'] * 100:.0f} / " + (f"{c['nodes'] / 1000:.1f}k" if c["nodes"] >= 10000 else f"{c['nodes']:,.0f}"))
        rows.append([short.get(g, g)] + cells)
    return rows


def equation(text, number):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(4), Pt(4)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(122.9), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(245.8), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run("\t" + text)
    r.italic, r.font.name, r.font.size = True, FONT, Pt(10)
    r2 = p.add_run("\t" + number)
    r2.font.name, r2.font.size = FONT, Pt(10)


sec_no, sub_no, fig_no, tab_no = 0, 0, 0, 0
in_refs, i = False, 0
body = lines[lines.index(next(l for l in lines if l.startswith("## Abstract"))):]
while i < len(body):
    ln = body[i]
    if not ln.strip():
        i += 1
        continue
    if ln.startswith("## "):
        head = ln[3:].strip()
        if head == "Abstract":
            i += 1
            while i < len(body) and not body[i].strip():
                i += 1
            p = para(align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, size=9, before=2, after=6)
            r = p.add_run("Abstract")
            r.bold, r.italic, r.font.size, r.font.name = True, True, Pt(9), FONT
            r = p.add_run(" - " + body[i].strip())
            r.bold, r.font.size, r.font.name = True, Pt(9), FONT
            i += 1
            continue
        in_refs = head == "References"
        if in_refs:
            p = para(align=WD_ALIGN_PARAGRAPH.CENTER, before=12, after=6, keep_next=True)
            r = p.add_run("References")
            r.font.small_caps, r.font.size, r.font.name = True, Pt(10), FONT
        else:
            sec_no += 1
            sub_no = 0
            p = para(align=WD_ALIGN_PARAGRAPH.CENTER, before=12, after=6, keep_next=True)
            a = p.add_run(f"{ROMAN[sec_no - 1]}. ")
            a.font.size, a.font.name = Pt(10), FONT
            b = p.add_run(head)
            b.font.small_caps, b.font.size, b.font.name = True, Pt(10), FONT
        i += 1
        continue
    if ln.startswith("### "):
        sub_no += 1
        p = para(align=WD_ALIGN_PARAGRAPH.LEFT, before=8, after=4, keep_next=True)
        r = p.add_run(f"{chr(64 + sub_no)}. {ln[4:].strip()}")
        r.italic, r.font.size, r.font.name = True, Pt(10), FONT
        i += 1
        continue
    if ln.startswith("Keywords:"):
        p = para(align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, size=9, before=2, after=6)
        r = p.add_run("Keywords - ")
        r.bold = r.italic = True
        r.font.size, r.font.name = Pt(9), FONT
        r = p.add_run(ln[len("Keywords:"):].strip())
        r.bold = r.italic = True
        r.font.size, r.font.name = Pt(9), FONT
        i += 1
        continue
    if ln.startswith("EQ: "):
        text, number = ln[4:].split(" | ")
        equation(text, number)
        i += 1
        continue
    if ln.startswith("TABLE: "):
        tab_no += 1
        cap = para(align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=0, keep_next=True)
        r = cap.add_run(f"TABLE {tab_no}")
        r.font.size, r.font.name = Pt(10), FONT
        t2 = para(align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=4, keep_next=True)
        r = t2.add_run(ln[7:].strip().upper())
        r.font.size, r.font.name = Pt(10), FONT
        i += 1
        if body[i].strip() == "@search_table":
            add_table(search_rows())
            i += 1
        else:
            rows = []
            while i < len(body) and body[i].startswith("|"):
                if not re.fullmatch(r"\|[\s\-|:]+\|", body[i].strip()):
                    rows.append(body[i].strip().strip("|").split("|"))
                i += 1
            add_table(rows)
        continue
    if ln.startswith("!["):
        fig_no += 1
        m = re.match(r"!\[(.*)\]\((.*)\)", ln)
        pic = doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.paragraph_format.space_before = Pt(4)
        pic.paragraph_format.keep_with_next = True
        pic.add_run().add_picture(str(HERE / m.group(2)), width=Pt(240))
        c = para(align=WD_ALIGN_PARAGRAPH.CENTER, size=9, before=2, after=8)
        r = c.add_run(m.group(1).strip())
        r.font.size, r.font.name = Pt(9), FONT
        i += 1
        continue
    if ln.startswith("> "):
        i += 1  # builder note, not part of the paper
        continue
    if in_refs and ln.startswith("["):
        p = para(align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        runs(p, ln, size=10, ital=True)
    else:
        p = para(align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True)
        runs(p, ln)
    i += 1

out = HERE / "paper.docx"
doc.save(out)
print("wrote", out)
